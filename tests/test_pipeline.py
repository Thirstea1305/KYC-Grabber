"""End-to-end tests: the always-on service detects mail and produces the workbook."""

from __future__ import annotations

import asyncio
import json
from pathlib import Path

from openpyxl import load_workbook

from kyc_grabber.mail.composer import parse_rfc822
from kyc_grabber.runtime import AppContext, build_context
from kyc_grabber.samples import SAMPLE_CSV, build_sample_eml
from tests.helpers import make_mail, make_settings, wait_for


# --------------------------------------------------------------------------- watcher


def test_watcher_detects_dropped_email_and_runs_the_full_pipeline(tmp_path: Path) -> None:
    async def scenario() -> None:
        context: AppContext = build_context(make_settings(tmp_path))
        await context.worker.start()
        await context.source.start()
        try:
            (context.settings.inbox_path / "request.eml").write_bytes(
                build_sample_eml(message_id="<watcher-1@corp.com>")
            )

            job = await wait_for(lambda: context.store.list_jobs(limit=1) or None)
            job_id = job[0]["id"]

            completed = await wait_for(
                lambda: (context.store.get_job(job_id) or {}).get("status") == "completed" and job_id,
                timeout=20.0,
            )
            assert completed

            stored = context.store.get_job(job_id)
            assert stored["total_codes"] == 6
            assert stored["processed_codes"] == 6
            assert stored["delivery_mode"] == "filesystem"

            workbook_path = Path(stored["workbook_path"])
            assert workbook_path.is_file()
            names = load_workbook(workbook_path).sheetnames
            assert names[0] == "Summary"
            assert set(names[1:]) == {"TP-0001", "TP-0002", "TP-0003", "TP-0007", "EXT-4242", "XX-9999"}

            outbox = sorted(context.settings.outbox_path.glob("reply-*.eml"))
            assert len(outbox) == 1
            reply = parse_rfc822(outbox[0].read_bytes())
            assert reply.subject.startswith("KYC report for")
            assert [a.filename for a in reply.attachments] == [workbook_path.name]

            # the source archived the request so it will not be processed twice
            assert not list(context.settings.inbox_path.glob("*.eml"))
        finally:
            await context.source.stop()
            await context.worker.stop()
            context.close()

    asyncio.run(scenario())


def test_processed_mail_is_archived_and_not_reprocessed(tmp_path: Path) -> None:
    async def scenario() -> None:
        context = build_context(make_settings(tmp_path))
        await context.source.start()
        try:
            inbox = context.settings.inbox_path
            (inbox / "one.eml").write_bytes(build_sample_eml(message_id="<archive-1@corp.com>"))
            await wait_for(lambda: context.store.list_jobs(limit=1))
            await asyncio.sleep(0.2)
            assert len(context.store.list_jobs(limit=10)) == 1
            assert len(list(context.settings.processed_path.glob("*.eml"))) == 1
        finally:
            await context.source.stop()
            context.close()

    asyncio.run(scenario())


# ------------------------------------------------------------------------ pipeline


def test_duplicate_message_id_is_ignored(tmp_path: Path) -> None:
    async def scenario() -> None:
        context = build_context(make_settings(tmp_path))
        try:
            first = await context.pipeline.ingest(make_mail(message_id="<dup@corp.com>"))
            second = await context.pipeline.ingest(make_mail(message_id="<dup@corp.com>"))
            assert first is not None
            assert second is None
            assert len([j for j in context.store.list_jobs() if j["status"] != "rejected"]) == 1
        finally:
            context.close()

    asyncio.run(scenario())


def test_email_without_csv_attachment_is_rejected_with_a_reason(tmp_path: Path) -> None:
    async def scenario() -> None:
        context = build_context(make_settings(tmp_path))
        try:
            mail = make_mail(filename="notes.txt", message_id="<nofile@corp.com>")
            assert await context.pipeline.ingest(mail) is None
            job = context.store.list_jobs(limit=1)[0]
            assert job["status"] == "rejected"
            assert "expected .csv" in job["reject_reason"]
        finally:
            context.close()

    asyncio.run(scenario())


def test_csv_without_valid_codes_is_rejected(tmp_path: Path) -> None:
    async def scenario() -> None:
        context = build_context(make_settings(tmp_path))
        try:
            mail = make_mail(csv_text="third_party_code\nbad code!!\n", message_id="<bad@corp.com>")
            assert await context.pipeline.ingest(mail) is None
            assert "No valid third-party codes" in context.store.list_jobs(limit=1)[0]["reject_reason"]
        finally:
            context.close()

    asyncio.run(scenario())


def test_code_limit_per_email_is_enforced(tmp_path: Path) -> None:
    async def scenario() -> None:
        settings = make_settings(tmp_path)
        settings.filter.max_codes_per_request = 3
        context = build_context(settings)
        try:
            csv_text = "third_party_code\n" + "\n".join(f"TP-{n:04d}" for n in range(1, 6)) + "\n"
            assert await context.pipeline.ingest(make_mail(csv_text=csv_text, message_id="<big@corp.com>")) is None
            assert "limit per email is 3" in context.store.list_jobs(limit=1)[0]["reject_reason"]
        finally:
            context.close()

    asyncio.run(scenario())


def test_external_outage_marks_items_failed_but_still_delivers(tmp_path: Path) -> None:
    async def scenario() -> None:
        settings = make_settings(tmp_path)
        settings.external.failure_rate = 1.0
        settings.external.max_retries = 1
        context = build_context(settings)
        try:
            job_id = await context.pipeline.ingest(make_mail(
                csv_text="third_party_code\nTP-0001\nTP-0002\n", message_id="<outage@corp.com>"
            ))
            assert job_id
            await context.pipeline.process_job(job_id)

            job = context.store.get_job(job_id)
            assert job["status"] == "completed"
            assert job["failed_codes"] == 2
            items = context.store.list_items(job_id)
            assert all(item["status"] == "failed" for item in items)
            assert all(item["match_status"] == "error" for item in items)
            assert all("unavailable" in (item["error"] or "") for item in items)
            assert Path(job["workbook_path"]).is_file()
        finally:
            context.close()

    asyncio.run(scenario())


def test_accepted_job_is_queued_and_rerun_reprocesses_it(tmp_path: Path) -> None:
    async def scenario() -> None:
        context = build_context(make_settings(tmp_path))
        await context.worker.start()
        try:
            job_id = await context.pipeline.ingest(make_mail(
                csv_text="third_party_code\nTP-0001\n", message_id="<rerun@corp.com>"
            ))
            assert job_id

            # ingest() hands the job to the worker pool, so no explicit enqueue is needed
            await wait_for(lambda: (context.store.get_job(job_id) or {}).get("status") == "completed")
            assert context.worker.status()["processed"] == 1

            context.store.update_job(job_id, status="queued")
            await context.worker.enqueue(job_id)
            await wait_for(lambda: context.worker.status()["processed"] == 2)
            assert (context.store.get_job(job_id) or {})["status"] == "completed"
            assert len(context.store.list_items(job_id)) == 1  # upserted, not duplicated
        finally:
            await context.worker.stop()
            context.close()

    asyncio.run(scenario())
