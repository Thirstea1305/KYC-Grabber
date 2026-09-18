from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from kyc_grabber.config import ExcelSettings
from kyc_grabber.excel_report import sanitize_sheet_name, write_workbook
from kyc_grabber.models import (
    ExternalRecord,
    InternalRecord,
    ItemResult,
    ItemStatus,
    MatchStatus,
    RiskRating,
)


def make_item(code: str, *, match=MatchStatus.MATCH, risk=RiskRating.LOW, internal=True, external=True,
              error: str | None = None) -> ItemResult:
    return ItemResult(
        code=code,
        status=ItemStatus.FAILED if error else ItemStatus.COMPLETED,
        match_status=match,
        risk_rating=risk,
        internal=InternalRecord(
            code=code, legal_name=f"{code} Internal Ltd", entity_type="Private Limited Company",
            country="GB", risk_rating=RiskRating.LOW, sanctions_hits=0, pep_flags=0,
        ) if internal else None,
        external=ExternalRecord(
            code=code, registry_name="Companies House (UK)", registration_number="A1",
            legal_name=f"{code} External Ltd", entity_type="Private Limited Company", country="GB",
            directors=["A. Ruiz - Director"], shareholders=["Holder Ltd (100%)"],
            adverse_media_count=1, adverse_media=["Local press coverage (2019)"],
        ) if external else None,
        differences=["Country: internal='GB' vs external='IE'"] if match is MatchStatus.MISMATCH else [],
        warnings=["check manually"] if match is MatchStatus.NOT_FOUND else [],
        error=error,
        duration_ms=42,
    )


def test_sanitize_sheet_name_enforces_excel_limits() -> None:
    used: set[str] = set()
    assert sanitize_sheet_name("TP-0001", used) == "TP-0001"

    long_name = "A" * 60
    assert len(sanitize_sheet_name(long_name, used)) <= 31

    forbidden = sanitize_sheet_name("AA/BB:CC*DD?EE[FF]GG", used)
    assert not any(char in forbidden for char in "[]:*?/\\")

    used_with_collision: set[str] = {"tp-0001"}
    assert sanitize_sheet_name("TP-0001", used_with_collision) == "TP-0001~2"
    assert sanitize_sheet_name("TP-0001", used_with_collision) == "TP-0001~3"


def test_duplicate_codes_get_unique_sheet_names(tmp_path: Path) -> None:
    items = [make_item("TP-0001"), make_item("TP-0001"), make_item("TP-0001")]
    path = write_workbook(
        job_id="J-TEST-1", sender="a@b.com", subject="[KYC] x", source_file="codes.csv",
        items=items, settings=ExcelSettings(output_dir=tmp_path, include_raw_payload=False),
        output_dir=tmp_path,
    )
    workbook = load_workbook(path)
    assert len(workbook.sheetnames) == len(set(workbook.sheetnames)) == 4  # summary + 3 sheets


def test_one_sheet_per_third_party_plus_summary(tmp_path: Path) -> None:
    items = [
        make_item("TP-0001"),
        make_item("TP-0002", match=MatchStatus.MISMATCH, risk=RiskRating.HIGH),
        make_item("TP-0003", match=MatchStatus.EXTERNAL_ONLY, internal=False, external=True),
        make_item("TP-0004", match=MatchStatus.NOT_FOUND, internal=False, external=False),
        make_item("TP-0005", match=MatchStatus.ERROR, error="external lookup timed out"),
    ]
    path = write_workbook(
        job_id="J-TEST-2", sender="a@b.com", subject="[KYC] x", source_file="codes.csv",
        items=items, settings=ExcelSettings(output_dir=tmp_path, include_raw_payload=True),
        output_dir=tmp_path,
    )

    workbook = load_workbook(path)
    assert workbook.sheetnames[0] == "Summary"
    assert set(workbook.sheetnames[1:]) == {"TP-0001", "TP-0002", "TP-0003", "TP-0004", "TP-0005"}

    summary = workbook["Summary"]
    values = [cell.value for row in summary.iter_rows() for cell in row]
    assert "J-TEST-2" in str(values)
    assert all(code in str(values) for code in ["TP-0001", "TP-0005"])

    detail = workbook["TP-0002"]
    detail_values = [cell.value for row in detail.iter_rows() for cell in row]
    assert any("MISMATCH" == str(v).upper() for v in detail_values)
    assert any("Country: internal='GB' vs external='IE'" in str(v) for v in detail_values)
    assert any("A. Ruiz - Director" in str(v) for v in detail_values)

    assert "Summary" in workbook.sheetnames


def test_summary_sheet_can_be_disabled(tmp_path: Path) -> None:
    path = write_workbook(
        job_id="J-TEST-3", sender="a@b.com", subject="s", source_file="c.csv",
        items=[make_item("TP-0001")],
        settings=ExcelSettings(output_dir=tmp_path, include_summary_sheet=False, include_raw_payload=False),
        output_dir=tmp_path,
    )
    assert load_workbook(path).sheetnames == ["TP-0001"]


def test_workbook_is_written_where_asked(tmp_path: Path) -> None:
    path = write_workbook(
        job_id="J-TEST-4", sender="a@b.com", subject="s", source_file="c.csv",
        items=[make_item("TP-0001")],
        settings=ExcelSettings(output_dir=tmp_path),
        output_dir=tmp_path / "nested",
    )
    assert path.exists()
    assert path.parent == tmp_path / "nested"
    assert path.suffix == ".xlsx"
