"""Excel report writer: one worksheet per third party, plus a summary cover sheet."""

from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from .config import ExcelSettings
from .models import ItemResult, MatchStatus, RiskRating

# ------------------------------------------------------------------ styling
TITLE_FILL = PatternFill("solid", fgColor="1F3864")
TITLE_FONT = Font(bold=True, size=14, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="D9E2F3")
SECTION_FONT = Font(bold=True, size=11, color="1F3864")
LABEL_FONT = Font(bold=True)
MUTED_FONT = Font(italic=True, size=9, color="595959")

STATUS_FILLS = {
    MatchStatus.MATCH: PatternFill("solid", fgColor="C6EFCE"),
    MatchStatus.MISMATCH: PatternFill("solid", fgColor="FFEB9C"),
    MatchStatus.INTERNAL_ONLY: PatternFill("solid", fgColor="DDEBF7"),
    MatchStatus.EXTERNAL_ONLY: PatternFill("solid", fgColor="DDEBF7"),
    MatchStatus.NOT_FOUND: PatternFill("solid", fgColor="E7E6E6"),
    MatchStatus.ERROR: PatternFill("solid", fgColor="FFC7CE"),
}

RISK_FILLS = {
    RiskRating.CRITICAL: PatternFill("solid", fgColor="FFC7CE"),
    RiskRating.HIGH: PatternFill("solid", fgColor="FFD9B3"),
    RiskRating.MEDIUM: PatternFill("solid", fgColor="FFEB9C"),
    RiskRating.LOW: PatternFill("solid", fgColor="C6EFCE"),
    RiskRating.UNRATED: PatternFill("solid", fgColor="F2F2F2"),
}

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")

INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
MAX_SHEET_NAME = 31


def sanitize_sheet_name(raw: str, used: set[str]) -> str:
    """Excel-safe, unique worksheet title derived from a third-party code."""
    name = INVALID_SHEET_CHARS.sub("-", (raw or "UNKNOWN")).strip().strip("'")
    name = name or "UNKNOWN"
    name = name[: MAX_SHEET_NAME - 4]
    candidate = name
    counter = 2
    while candidate.lower() in used:
        suffix = f"~{counter}"
        candidate = f"{name[: MAX_SHEET_NAME - len(suffix)]}{suffix}"
        counter += 1
    used.add(candidate.lower())
    return candidate


def _title(ws: Worksheet, text: str, *, span: int = 5, subtitle: str | None = None) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    cell = ws.cell(row=1, column=1, value=text)
    cell.fill = TITLE_FILL
    cell.font = TITLE_FONT
    cell.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 24
    row = 2
    if subtitle:
        ws.cell(row=row, column=1, value=subtitle).font = MUTED_FONT
        row += 1
    return row + 1


def _section(ws: Worksheet, row: int, text: str, *, span: int = 5) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.fill = SECTION_FILL
    cell.font = SECTION_FONT
    return row + 1


def _kv(ws: Worksheet, row: int, label: str, value, *, fill=None, bold_label: bool = True) -> int:
    label_cell = ws.cell(row=row, column=1, value=label)
    label_cell.font = LABEL_FONT if bold_label else Font()
    label_cell.alignment = TOP
    value_cell = ws.cell(row=row, column=2, value=value if value not in (None, "") else "-")
    value_cell.alignment = WRAP
    if fill is not None:
        value_cell.fill = fill
    return row + 1


def _table_header(ws: Worksheet, row: int, headers: list[str], *, start_column: int = 1) -> int:
    for offset, header in enumerate(headers):
        cell = ws.cell(row=row, column=start_column + offset, value=header)
        cell.font = LABEL_FONT
        cell.fill = SECTION_FILL
        cell.border = BORDER
    return row + 1


def _table_row(ws: Worksheet, row: int, values: list[str], *, start_column: int = 1) -> int:
    for offset, value in enumerate(values):
        cell = ws.cell(row=row, column=start_column + offset, value=value)
        cell.alignment = WRAP
        cell.border = BORDER
    return row + 1


def _set_widths(ws: Worksheet, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def _item_view(item: ItemResult) -> dict:
    """Flatten an ItemResult into the plain structure the report renders."""
    return {
        "code": item.code,
        "match_status": (item.match_status or MatchStatus.ERROR).value if item.match_status else "-",
        "risk": item.risk_rating.value,
        "internal": item.internal.model_dump(mode="json") if item.internal else None,
        "external": item.external.model_dump(mode="json") if item.external else None,
        "differences": list(item.differences),
        "warnings": list(item.warnings),
        "error": item.error,
        "status": item.status.value,
        "duration_ms": item.duration_ms,
    }


def build_workbook(
    *,
    job_id: str,
    sender: str,
    subject: str,
    source_file: str,
    items: list[ItemResult],
    settings: ExcelSettings,
    generated_at: datetime | None = None,
) -> Workbook:
    """Build the multi-sheet workbook returned to the requester."""
    generated_at = generated_at or datetime.now(timezone.utc)
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names: set[str] = set()

    if settings.include_summary_sheet:
        _build_summary_sheet(workbook, job_id, sender, subject, source_file, items, generated_at, used_names)

    for item in items:
        _build_detail_sheet(workbook, item, job_id, generated_at, used_names, settings)

    if not workbook.worksheets:  # extremely defensive: openpyxl requires >=1 sheet
        workbook.create_sheet("Empty")

    return workbook


def _build_summary_sheet(
    workbook: Workbook,
    job_id: str,
    sender: str,
    subject: str,
    source_file: str,
    items: list[ItemResult],
    generated_at: datetime,
    used_names: set[str],
) -> None:
    ws = workbook.create_sheet(sanitize_sheet_name("Summary", used_names))
    _set_widths(ws, {"A": 22, "B": 40, "C": 34, "D": 10, "E": 16, "F": 12, "G": 10, "H": 10, "I": 45})

    row = _title(ws, f"KYC Grabber - report {job_id}", span=9,
                 subtitle="Automated output. Sources marked 'simulated' contain synthetic demo data only.")
    row = _kv(ws, row, "Requested by", sender, bold_label=False)
    row = _kv(ws, row, "Email subject", subject, bold_label=False)
    row = _kv(ws, row, "Source file", source_file, bold_label=False)
    row = _kv(ws, row, "Generated (UTC)", generated_at.strftime("%Y-%m-%d %H:%M:%S"), bold_label=False)
    row = _kv(ws, row, "Third parties", len(items), bold_label=False)
    total_ms = sum(item.duration_ms for item in items)
    row = _kv(ws, row, "Total enrichment time", f"{total_ms / 1000:.1f}s", bold_label=False)
    row += 1

    row = _section(ws, row, "Findings", span=9)
    row = _table_header(ws, row, [
        "Third-party code", "External legal name", "Internal legal name", "Country",
        "Match status", "Risk rating", "Internal", "External", "Notes",
    ])

    order = {MatchStatus.MISMATCH: 0, MatchStatus.NOT_FOUND: 1, MatchStatus.EXTERNAL_ONLY: 2,
             MatchStatus.INTERNAL_ONLY: 3, MatchStatus.ERROR: 4, MatchStatus.MATCH: 5}
    for item in sorted(items, key=lambda entry: (order.get(entry.match_status or MatchStatus.ERROR, 9), entry.code)):
        internal = item.internal
        external = item.external
        notes = "; ".join(item.differences[:2] + item.warnings[:1]) or (item.error or "")
        status = item.match_status or MatchStatus.ERROR
        match_cell = ws.cell(row=row, column=5, value=status.value)
        match_cell.fill = STATUS_FILLS.get(status, PatternFill())
        risk_cell = ws.cell(row=row, column=6, value=item.risk_rating.value)
        risk_cell.fill = RISK_FILLS.get(item.risk_rating, PatternFill())

        _table_row(ws, row, [
            item.code,
            (external.legal_name if external else "-"),
            (internal.legal_name if internal else "-"),
            (external.country if external else (internal.country if internal else "-")),
            "", "",
            "yes" if internal else "no",
            "yes" if external else "no",
            notes,
        ])
        # re-apply the coloured cells overwritten by _table_row
        ws.cell(row=row, column=5).fill = STATUS_FILLS.get(status, PatternFill())
        ws.cell(row=row, column=5).value = status.value
        ws.cell(row=row, column=6).fill = RISK_FILLS.get(item.risk_rating, PatternFill())
        ws.cell(row=row, column=6).value = item.risk_rating.value
        row += 1

    ws.freeze_panes = "A11"


def _build_detail_sheet(
    workbook: Workbook,
    item: ItemResult,
    job_id: str,
    generated_at: datetime,
    used_names: set[str],
    settings: ExcelSettings,
) -> None:
    view = _item_view(item)
    ws = workbook.create_sheet(sanitize_sheet_name(item.code, used_names))
    _set_widths(ws, {"A": 30, "B": 78, "C": 34, "D": 34})

    status = item.match_status or MatchStatus.ERROR
    row = _title(ws, f"Third party: {item.code}", span=4,
                 subtitle=f"Job {job_id} | generated {generated_at.strftime('%Y-%m-%d %H:%M:%S')} UTC")

    row = _section(ws, row, "1. Reconciliation outcome", span=4)
    row = _kv(ws, row, "Match status", status.value, fill=STATUS_FILLS.get(status))
    row = _kv(ws, row, "Overall risk rating", item.risk_rating.value, fill=RISK_FILLS.get(item.risk_rating))
    row = _kv(ws, row, "Internal record found", "yes" if item.internal else "no")
    row = _kv(ws, row, "External record found", "yes" if item.external else "no")
    row = _kv(ws, row, "Enrichment time", f"{item.duration_ms} ms")
    row += 1

    row = _section(ws, row, "2. Internal database record", span=4)
    internal = view["internal"]
    if internal:
        row = _table_header(ws, row, ["Field", "Value"], )
        for label, key in [
            ("Third-party code", "code"), ("Legal name", "legal_name"), ("Entity type", "entity_type"),
            ("Country", "country"), ("Risk rating (declared)", "risk_rating"), ("Status", "status"),
            ("Business unit", "business_unit"), ("Onboarded at", "onboarded_at"),
            ("Sanctions hits", "sanctions_hits"), ("PEP flags", "pep_flags"), ("Notes", "notes"),
        ]:
            row = _table_row(ws, row, [label, str(internal.get(key) if internal.get(key) is not None else "-")])
        row = _table_row(ws, row, ["Source", internal.get("source", "-")])
    else:
        row = _table_row(ws, row, ["No internal record returned for this code.", "-"])
    row += 1

    row = _section(ws, row, "3. External database record", span=4)
    external = view["external"]
    if external:
        row = _table_header(ws, row, ["Field", "Value"])
        for label, key in [
            ("Registry", "registry_name"), ("Registration number", "registration_number"),
            ("Legal name", "legal_name"), ("Entity type", "entity_type"), ("Country", "country"),
            ("Incorporation date", "incorporation_date"), ("Status", "status"),
            ("Registered address", "registered_address"),
            ("Sanctions listed", "sanctions_listed"), ("Sanctions lists", "sanctions_lists"),
            ("PEP match", "pep_match"), ("Adverse media count", "adverse_media_count"),
            ("Source URL", "source_url"), ("Fetched at", "fetched_at"), ("Data source", "source"),
        ]:
            value = external.get(key)
            if isinstance(value, list):
                value = ", ".join(str(entry) for entry in value) or "-"
            if isinstance(value, bool):
                value = "yes" if value else "no"
            row = _table_row(ws, row, [label, str(value if value is not None else "-")])

        row += 1
        row = _section(ws, row, "3a. Directors", span=4)
        if external.get("directors"):
            row = _table_header(ws, row, ["Name", "Role"])
            for entry in external["directors"]:
                name, _, role = str(entry).partition(" - ")
                row = _table_row(ws, row, [name, role or "-"])
        else:
            row = _table_row(ws, row, ["-", "-"])

        row += 1
        row = _section(ws, row, "3b. Shareholders", span=4)
        if external.get("shareholders"):
            row = _table_header(ws, row, ["Holder", "Details"])
            for entry in external["shareholders"]:
                holder, _, details = str(entry).partition(" (")
                row = _table_row(ws, row, [holder, f"({details}" if details else "-"])
        else:
            row = _table_row(ws, row, ["-", "-"])

        row += 1
        row = _section(ws, row, "3c. Adverse media", span=4)
        if external.get("adverse_media"):
            for entry in external["adverse_media"]:
                row = _table_row(ws, row, ["-", str(entry)])
        else:
            row = _table_row(ws, row, ["-", "No adverse media found."])
    else:
        row = _table_row(ws, row, ["No external record returned for this code.", "-"])
    row += 1

    row = _section(ws, row, "4. Differences and warnings", span=4)
    if item.differences:
        row = _table_header(ws, row, ["Difference", "Detail"])
        for entry in item.differences:
            label, _, _detail = entry.partition(": ")
            row = _table_row(ws, row, [label or "Difference", entry])
    if item.warnings:
        if not item.differences:
            row = _table_header(ws, row, ["Warning", "Detail"])
        for entry in item.warnings:
            row = _table_row(ws, row, ["Warning", entry])
    if item.error:
        row = _table_row(ws, row, ["Error", item.error])
    if not (item.differences or item.warnings or item.error):
        row = _table_row(ws, row, ["No differences detected.", "-"])
    row += 1

    if settings.include_raw_payload:
        row = _section(ws, row, "5. Raw source payloads (as received)", span=4)
        payload = json.dumps({"internal": view["internal"], "external": view["external"]}, indent=2, default=str)
        cell = ws.cell(row=row, column=1, value=payload)
        cell.alignment = WRAP
        cell.font = Font(name="Consolas", size=9)
        ws.merge_cells(start_row=row, start_column=1, end_row=row + 40, end_column=4)
        row += 42

    ws.cell(row=row, column=1,
            value="Disclaimer: data labelled 'simulated' is synthetic and must not be used for real decisions.").font = MUTED_FONT


def default_filename(job_id: str, when: datetime | None = None) -> str:
    when = when or datetime.now(timezone.utc)
    return f"kyc-report-{job_id}-{when.strftime('%Y%m%d-%H%M%S')}.xlsx"


def write_workbook(
    *,
    job_id: str,
    sender: str,
    subject: str,
    source_file: str,
    items: list[ItemResult],
    settings: ExcelSettings,
    output_dir: Path,
    generated_at: datetime | None = None,
) -> Path:
    """Render and persist the workbook, returning the written path."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook = build_workbook(
        job_id=job_id,
        sender=sender,
        subject=subject,
        source_file=source_file,
        items=items,
        settings=settings,
        generated_at=generated_at,
    )
    path = output_dir / default_filename(job_id, generated_at)
    workbook.save(path)
    return path
